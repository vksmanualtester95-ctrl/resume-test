#!/usr/bin/env python3
"""
check.py - Shows you EXACTLY what an ATS reads from your resume, then audits it.

Zero heavy dependencies. Uses pypdf (~1 MB), not poppler (~100 MB).

Usage:
  python3 scripts/check.py input/Vicky_Kumar-Azure_Data_Engineer.pdf
  python3 scripts/check.py output/resume.pdf --text     # dump raw text only
"""
import argparse, os, re, sys, zipfile
from collections import Counter

OK, WARN, FAIL = "  ok  ", " WARN ", " FAIL "

STD_HEADINGS = {"summary", "profile", "objective", "experience", "work experience",
                "professional experience", "employment", "education", "skills",
                "technical skills", "core skills", "projects", "certifications",
                "certificates", "awards", "publications", "languages", "achievements"}

AI_WORDS = ["spearheaded", "leveraged", "orchestrated", "championed", "delve",
            "tapestry", "seamless", "synergy", "results-driven", "detail-oriented",
            "passionate about", "proven track record", "team player", "self-starter",
            "cutting-edge", "state-of-the-art", "robust solutions", "holistic",
            "thought leader", "dynamic professional", "pivotal role"]

WEAK = ["responsible for", "worked on", "helped with", "assisted with",
        "duties included", "tasked with", "involved in", "participated in"]

ICON_FONTS = ["fontawesome", "wingdings", "webdings", "glyphicons", "materialicons",
              "ionicons", "feather", "zapfdingbats", "symbol"]


def read_pdf(path):
    try:
        import pypdf
    except ImportError:
        sys.exit("Run:  pip install pypdf")
    r = pypdf.PdfReader(path)
    text = "\n".join(p.extract_text() or "" for p in r.pages)
    fonts = set()
    for p in r.pages:
        try:
            for _, v in p["/Resources"]["/Font"].items():
                fonts.add(str(v.get_object().get("/BaseFont", "")))
        except Exception:
            pass
    return text, {"pages": len(r.pages), "fonts": sorted(fonts)}


def read_docx(path):
    meta = {"tables": 0, "textboxes": 0, "images": 0}
    try:
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        meta["tables"] = len(d.tables)
        for t in d.tables:
            for row in t.rows:
                for c in row.cells:
                    parts.append(c.text)
        meta["images"] = len(d.inline_shapes)
        for s in d.sections:
            for p in list(s.header.paragraphs) + list(s.footer.paragraphs):
                if p.text.strip():
                    meta["headers_footers"] = meta.get("headers_footers", 0) + 1
        return "\n".join(parts), meta
    except ImportError:
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml").decode("utf-8", "replace")
        meta["textboxes"] = xml.count("<w:txbxContent")
        meta["tables"] = xml.count("<w:tbl>")
        return "\n".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml)), meta


def read_any(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".docx":
        return read_docx(path)
    return open(path, encoding="utf-8", errors="replace").read(), {}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--text", action="store_true", help="print raw text and exit")
    a = ap.parse_args()
    if not os.path.exists(a.file):
        sys.exit(f"not found: {a.file}")

    text, meta = read_any(a.file)

    if a.text:
        print(text)
        return 0

    rows, score = [], 100

    def add(level, msg, fix=""):
        nonlocal score
        if level == FAIL:
            score -= 14
        elif level == WARN:
            score -= 5
        rows.append((level, msg, fix))

    print("=" * 72)
    print(" WHAT THE ATS ACTUALLY READS (first 6 lines)")
    print("=" * 72)
    for line in [l for l in text.splitlines() if l.strip()][:6]:
        print("  " + repr(line)[:150])
    print()

    # --- text layer -------------------------------------------------------
    words = re.findall(r"[A-Za-z]{2,}", text)
    if len(words) < 80:
        add(FAIL, f"Almost no text extracted ({len(words)} words)",
            "PDF is image-based or text is outlined. Re-export as a real text PDF.")
        for l, m, f in rows:
            print(f"[{l}] {m}")
        return 2
    add(OK, f"Text layer present ({len(words)} words)")

    # --- icon fonts -------------------------------------------------------
    bad_fonts = [f for f in meta.get("fonts", [])
                 if any(k in f.lower() for k in ICON_FONTS)]
    if bad_fonts:
        add(FAIL, "Icon font embedded: " + ", ".join(bad_fonts),
            "Icons extract as junk glyphs and glue your contact fields together. "
            "Replace with plain text labels: 'Phone:' 'Email:' 'LinkedIn:'")

    # --- broken kerning / letter splits -----------------------------------
    splits = sorted(set(re.findall(r"\b[A-Z]\s[a-z]{2,}\b", text)))
    if splits:
        add(FAIL, "Words split by the PDF: " + ", ".join(repr(s) for s in splits[:6]),
            "The text layer breaks these words. Keyword search will miss them. "
            "Regenerate the PDF (browser Print-to-PDF or Overleaf pdfLaTeX).")

    # --- mojibake ---------------------------------------------------------
    for pat, msg in [(r"[\uFB00-\uFB06]", "unresolved ff/fi/fl ligatures"),
                     (r"\uFFFD", "replacement characters (broken encoding)"),
                     (r"(?:[A-Za-z] ){6,}[A-Za-z]", "letter-spaced text")]:
        if re.search(pat, text):
            add(FAIL, f"Damaged extraction: {msg}",
                "Regenerate the PDF from a clean source.")

    # --- glued title / date ----------------------------------------------
    glued = [l for l in text.splitlines()
             if re.search(r"[a-z\)][A-Z][a-z]{2}\s+\d{4}", l)]
    if glued:
        add(FAIL, "Job title glued to its date: " + repr(glued[0][:70]),
            "Parsers read this as one field and lose the title. "
            "Put the date on its own line, or separate with ' | '.")

    # --- multi-column -----------------------------------------------------
    lines = [l for l in text.splitlines() if l.strip()]
    gutters = [l for l in lines if re.search(r"\S {6,}\S", l)]
    ratio = len(gutters) / max(len(lines), 1)
    if ratio > 0.35:
        add(FAIL, f"Multi-column layout ({ratio:.0%} of lines have a wide gap)",
            "Workday/Taleo/iCIMS flatten left-to-right and interleave your "
            "skills into your job titles. Go single column.")
    else:
        add(OK, f"Single-column reading order ({ratio:.0%} gutter lines)")

    # --- headings ---------------------------------------------------------
    found = [l.strip() for l in lines
             if re.sub(r"[^a-z ]", "", l.strip().lower()) in STD_HEADINGS
             and len(l.strip()) < 45]
    have = {re.sub(r"[^a-z ]", "", f.lower()).strip() for f in found}
    missing = [n for n in ("experience", "education", "skills")
               if not any(n in h for h in have)]
    if missing:
        add(FAIL, "Missing standard heading(s): " + ", ".join(missing),
            "Parsers anchor on the literal words Experience / Education / Skills.")
    else:
        add(OK, "Standard headings found: " + ", ".join(found[:7]))

    # --- contact ----------------------------------------------------------
    head = "\n".join(lines[:6])
    email = re.search(r"[\w.\-+]+@[\w.\-]+\.\w{2,}", head)
    phone = re.search(r"(?:\+\d{1,3}[\s\-]?)?\d{5}[\s\-]?\d{5}|\d{10}", head)
    if email:
        e = email.group()
        before = head[max(0, email.start() - 1):email.start()]
        if before and (before.isalnum() or ord(before) > 0x2000):
            add(FAIL, f"Email is glued to the previous field: ...{before}{e[:20]}",
                "Some parsers will grab the whole blob and store a broken email. "
                "Separate contact fields with ' | ' and drop the icons.")
        else:
            add(OK, f"Email clean in top block: {e}")
    else:
        add(FAIL, "No clean email in the top 6 lines",
            "Put it on line 2 or 3 as plain text.")
    add(OK if phone else FAIL, "Phone found" if phone else "No phone in top block")

    if not re.search(r"\b(India|Bengaluru|Bangalore|Mumbai|Pune|Delhi|Noida|Gurgaon|"
                     r"Hyderabad|Chennai|Kolkata|Bhubaneswar|Remote)\b", head, re.I):
        add(WARN, "No city/location in the contact block",
            "Recruiters filter by location and Workday has a location field. "
            "Add 'City, State, India' on the contact line.")
    else:
        add(OK, "Location present in contact block")

    # --- dates ------------------------------------------------------------
    mon = len(re.findall(r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
                         r"[a-z]*\.?\s+\d{4}\b", text))
    if mon == 0:
        add(FAIL, "No 'Mon YYYY' employment dates found",
            "Use 'Mar 2022 - Aug 2023'. Missing dates break tenure calculation.")
    else:
        add(OK, f"{mon} 'Mon YYYY' dates found")

    # --- length -----------------------------------------------------------
    pages = meta.get("pages")
    if pages:
        if pages > 2:
            add(FAIL, f"{pages} pages",
                "Workday clips past page 2. Cut it down.")
        elif pages == 2:
            add(WARN, "2 pages",
                "Under 8 years of experience, 1 page is the norm and gets read more.")
        else:
            add(OK, "1 page")

    # --- content quality --------------------------------------------------
    body = [l for l in lines if len(l) > 40]
    numeric = [l for l in body if re.search(r"\d", l)]
    dens = len(numeric) / max(len(body), 1)
    add(OK if dens >= 0.4 else WARN,
        f"Quantification: {dens:.0%} of content lines have a number",
        "" if dens >= 0.4 else "Target 50-70%. Scope counts: records, sources, "
                               "pipelines, users, team size.")

    low = text.lower()
    ai = [w for w in AI_WORDS if w in low]
    if ai:
        add(FAIL if len(ai) >= 5 else WARN, f"AI/buzzword tells: {', '.join(ai[:8])}",
            "49% of hiring managers dismiss resumes that read as AI-written. "
            "Replace with concrete, specific language.")
    else:
        add(OK, "No common AI buzzword tells")

    weak = [w for w in WEAK if w in low]
    if weak:
        add(WARN, f"Duty phrasing: {', '.join(weak[:5])}",
            "Rewrite as [verb] + [what you did] + [measured result].")

    rounds = re.findall(r"\b(?:10|20|25|30|40|50|60|70|75|80|90|100)%", text)
    if len(rounds) >= 4:
        add(WARN, f"Round-number stacking: {Counter(rounds).most_common(5)}",
            "Top recruiter tell for AI drafts. Use the real odd numbers "
            "(your 83% and 99.9% are exactly right - make the others match).")

    if body:
        lens = [len(l) for l in body]
        if max(lens) - min(lens) < 25 and len(body) > 6:
            add(WARN, "Bullets are near-identical in length",
                "Reads as machine-generated. Vary them on purpose.")

    # --- docx extras ------------------------------------------------------
    if meta.get("tables"):
        add(FAIL, f"{meta['tables']} table(s) in the DOCX",
            "Table cells land in the wrong parsed fields. Use paragraphs.")
    if meta.get("textboxes"):
        add(FAIL, f"{meta['textboxes']} text box(es)", "Most parsers never read them.")
    if meta.get("headers_footers"):
        add(WARN, "Content in header/footer",
            "Workday and Taleo suppress header/footer text.")
    if meta.get("images"):
        add(WARN, f"{meta['images']} image(s)", "No text layer. Remove photos/logos.")

    # --- filename ---------------------------------------------------------
    fn = os.path.basename(a.file)
    if re.search(r"[^\w.\-]", fn):
        add(WARN, f"Filename has odd characters: {fn}",
            "Use Vicky_Kumar_Azure_Data_Engineer.pdf")

    # --- report -----------------------------------------------------------
    score = max(0, score)
    fails = sum(1 for r in rows if r[0] == FAIL)
    warns = sum(1 for r in rows if r[0] == WARN)
    print("=" * 72)
    print(f" PARSE AUDIT   score {score}/100   {fails} fail   {warns} warn")
    print("=" * 72)
    order = {FAIL: 0, WARN: 1, OK: 2}
    for lvl, msg, fix in sorted(rows, key=lambda r: order[r[0]]):
        print(f"[{lvl}] {msg}")
        if fix:
            for chunk in re.findall(r".{1,64}(?:\s|$)", fix):
                print(f"         {chunk.strip()}")
    print("-" * 72)
    print(" Parse score = machine readability only. It does not predict interviews.")
    print(" Run with --text to read the full extraction yourself.")
    return 2 if fails else (1 if warns else 0)


if __name__ == "__main__":
    sys.exit(main())
